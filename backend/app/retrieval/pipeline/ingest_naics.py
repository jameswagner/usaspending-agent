"""Ingest the Census 2022 NAICS reference files into chunk records, for
semantic lookup ("what NAICS code is custom software development") - not
for search_guide, which stays scoped to conceptual/definitional questions.
This is the resolve_naics_code tool's own separate index (see #6, #30).

Three source files, three different roles - not redundant:
  - 2-6 digit_2022_Codes.xlsx: canonical {code: title} for every hierarchy
    level (2-6 digit). Used for the title, not Descriptions.xlsx's own Title
    column, which has a stray trailing "T" (trilateral-agreement marker) on
    some rows - sidestepped entirely by taking titles from this file instead.
  - 2022_NAICS_Descriptions.xlsx: {code: description}, the real prose text.
  - 2022_NAICS_Index_File.xlsx: {code: [synonym phrases]}, 6-digit only -
    the natural-language phrasing official titles don't cover (e.g.
    "software" as a business description, not just a title word).

Reuses the exact chunk shape ingest_glossary.py already established
(id/source/page_start/page_end/paragraph_index/text/term/slug/related_slugs)
so vector_index.py/bm25_index.py need zero changes - term=title, slug=code,
related_slugs unused. Pointed at its own Chroma collection/Whoosh dir (not
the Guide/Glossary one) via CHROMA_DB_DIR/WHOOSH_INDEX_DIR env overrides at
index-build time, so it never competes with search_guide's retrieval.

Usage:
  python -m backend.app.retrieval.pipeline.ingest_naics
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl

# One sector's description (31-33, Manufacturing) embeds an HTML table for a
# two-column examples list - stripped rather than indexed as literal markup.
_HTML_TAG_RE = re.compile(r"<[^>]+>")

SOURCE_NAME = "NAICS Codes"

CODES_TITLES_FILE = "data/raw/2-6 digit_2022_Codes.xlsx"
DESCRIPTIONS_FILE = "data/raw/2022_NAICS_Descriptions.xlsx"
INDEX_FILE = "data/raw/2022_NAICS_Index_File.xlsx"


@dataclass
class NAICSChunk:
    id: str
    source: str
    page_start: int
    page_end: int
    paragraph_index: int
    text: str
    term: str
    slug: str
    related_slugs: list[str]


def _code_str(value) -> str:
    # Most codes load as int/float via openpyxl - normalize to a plain
    # digit string (no leading zeros in real 2022 NAICS codes, so int()
    # round-tripping is safe). But some top-level sectors are combined
    # ranges stored as text (e.g. "31-33" for Manufacturing, spanning
    # sectors 31/32/33) - found live when this raised on that first real
    # value, not assumed - those already come through as str and pass
    # through unchanged.
    if isinstance(value, str):
        return value.strip()
    return str(int(value))


def load_titles(path: str) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    titles: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        # Row 2 is a blank separator row in this file (confirmed by
        # inspection, not assumed) - skip any row without a real code.
        code = row[1]
        title = row[2]
        if code is None or title is None:
            continue
        titles[_code_str(code)] = title.strip()
    wb.close()
    return titles


def load_descriptions(path: str) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    descriptions: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code, _title, description = row[0], row[1], row[2]
        if code is None or description is None:
            continue
        descriptions[_code_str(code)] = _HTML_TAG_RE.sub(" ", description).strip()
    wb.close()
    return descriptions


def load_synonyms(path: str) -> dict[str, list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    synonyms: dict[str, list[str]] = defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        code, phrase = row[0], row[1]
        if code is None or phrase is None:
            continue
        synonyms[_code_str(code)].append(phrase.strip())
    wb.close()
    return dict(synonyms)


def build_chunk(code: str, title: str, description: str, synonyms: list[str]) -> NAICSChunk:
    text_parts = [f"{code} - {title}"]
    if description:
        text_parts.append(description)
    if synonyms:
        text_parts.append("Also known as: " + "; ".join(synonyms))
    return NAICSChunk(
        id=f"NAICS_{code}",
        source=SOURCE_NAME,
        page_start=0,
        page_end=0,
        paragraph_index=0,
        text="\n\n".join(text_parts),
        term=title,
        slug=code,
        related_slugs=[],
    )


def ingest_naics_to_chunks(out_path: Path) -> list[NAICSChunk]:
    titles = load_titles(CODES_TITLES_FILE)
    descriptions = load_descriptions(DESCRIPTIONS_FILE)
    synonyms = load_synonyms(INDEX_FILE)

    chunks = [
        build_chunk(code, title, descriptions.get(code, ""), synonyms.get(code, []))
        for code, title in titles.items()
    ]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as fh:
        for c in chunks:
            fh.write(json.dumps(asdict(c), ensure_ascii=False) + "\n")

    return chunks


def main():
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="data/chunks/naics_chunks.jsonl")
    args = parser.parse_args()

    chunks = ingest_naics_to_chunks(Path(args.out))
    print(f"Wrote {len(chunks)} chunks to {args.out}")


if __name__ == "__main__":
    main()
