"""Ingest the GSA PSC (Product and Service Code) manual into chunk records,
for semantic lookup - the psc_code counterpart to ingest_naics.py. Same
role: resolve_psc_code (not yet wired up - see #6, #30) resolves a
plain-English description to a code for search_awards/get_spending_by_category's
psc_code filter, separate from search_guide's definitional corpus.

One source file, but it carries retired/historical duplicate rows - the
same PSC code can appear more than once across different date ranges with
different text (e.g. code 1005 has a current row, END DATE null, and a
historical row covering 1978-2011 with slightly different wording).
Filtering to END DATE is null gives exactly the active set with zero
duplicates - confirmed against the real file (2,540 rows), not assumed.

Reuses the same chunk shape as ingest_glossary.py/ingest_naics.py
(id/source/page_start/page_end/paragraph_index/text/term/slug/related_slugs)
so vector_index.py/bm25_index.py need no changes - term=name, slug=code.

Usage:
  python -m backend.app.retrieval.pipeline.ingest_psc
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl

SOURCE_NAME = "PSC Codes"

PSC_FILE = "data/raw/PSC April 2025.xlsx"

# Column indices in the source sheet (0-based) - named here since the
# sheet has 14 columns and index-only access elsewhere would be unreadable.
COL_CODE = 0
COL_NAME = 1
COL_END_DATE = 3
COL_FULL_NAME = 4
COL_INCLUDES = 5
COL_EXCLUDES = 6
COL_NOTES = 7


@dataclass
class PSCChunk:
    id: str
    source: str
    page_start: int
    page_end: int
    paragraph_index: int
    text: str
    term: str
    slug: str
    related_slugs: list[str]


def _code_str(value) -> str:
    # PSC codes are a mix of purely numeric (1005 -> "1005") and
    # alphanumeric (e.g. "7A", "7A20", "A", "R414") - openpyxl hands back
    # a float for the former, str for the latter (confirmed against the
    # real file, not assumed).
    if isinstance(value, str):
        return value.strip()
    return str(int(value))


def load_active_rows(path: str) -> list[tuple]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[COL_END_DATE] is None]
    wb.close()
    return rows


def build_chunk(row: tuple) -> PSCChunk:
    code = _code_str(row[COL_CODE])
    name = row[COL_NAME].strip()

    text_parts = [f"{code} - {name}"]
    full_name = row[COL_FULL_NAME]
    if full_name and full_name.strip() != name:
        text_parts.append(full_name.strip())
    includes = row[COL_INCLUDES]
    if includes:
        text_parts.append(f"Includes: {includes.strip()}")
    excludes = row[COL_EXCLUDES]
    if excludes:
        text_parts.append(f"Excludes: {excludes.strip()}")
    notes = row[COL_NOTES]
    if notes:
        text_parts.append(notes.strip())

    return PSCChunk(
        id=f"PSC_{code}",
        source=SOURCE_NAME,
        page_start=0,
        page_end=0,
        paragraph_index=0,
        text="\n\n".join(text_parts),
        term=name,
        slug=code,
        related_slugs=[],
    )


def ingest_psc_to_chunks(out_path: Path) -> list[PSCChunk]:
    rows = load_active_rows(PSC_FILE)
    chunks = [build_chunk(row) for row in rows]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as fh:
        for c in chunks:
            fh.write(json.dumps(asdict(c), ensure_ascii=False) + "\n")

    return chunks


def main():
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="data/chunks/psc_chunks.jsonl")
    args = parser.parse_args()

    chunks = ingest_psc_to_chunks(Path(args.out))
    print(f"Wrote {len(chunks)} chunks to {args.out}")


if __name__ == "__main__":
    main()
